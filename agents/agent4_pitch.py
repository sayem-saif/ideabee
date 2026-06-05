from __future__ import annotations

import json
import os
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

from providers import ProviderError, extract_json, hf_chat, openrouter_chat


def hex_to_rgb(hex_str: str) -> RGBColor:
    try:
        hex_str = hex_str.lstrip('#')
        if len(hex_str) == 3:
            hex_str = ''.join([c*2 for c in hex_str])
        r = int(hex_str[0:2], 16)
        g = int(hex_str[2:4], 16)
        b = int(hex_str[4:6], 16)
        return RGBColor(r, g, b)
    except Exception:
        return RGBColor(2, 132, 199) # Fallback to sky blue #0284c7


def _get_python_fallback_slides(master_profile: dict[str, Any]) -> list[dict[str, Any]]:
    startup_name = master_profile.get("startup_name", "the company")
    tagline = master_profile.get("tagline", "Transforming the industry.")
    blueprint = master_profile.get("asset_blueprint", {})
    brand = master_profile.get("brand_identity", {})

    # Competitors
    competitors = master_profile.get("competitors", [])
    competitors_str = ", ".join(competitors) if competitors else "Competitor analysis unavailable"

    # TAM/SAM/SOM or Claims
    claims_list = master_profile.get("research_claims", [])
    claims_strs = [c["statement"] for c in claims_list if c.get("used_in_assets")]
    market_opp_str = master_profile.get("market_opportunity", "Market details ready.")
    if not claims_strs:
        claims_strs = ["Market opportunity estimates", "Market data unavailable"]

    # Team & Financials
    team_val = master_profile.get("team", "To Be Determined")
    financials_val = master_profile.get("financials", "Not Provided")

    return [
        {
            "title": startup_name,
            "bullets": [
                tagline,
                blueprint.get("hero_message", "Innovation Simplified"),
                "Widescreen investor deck presentation."
            ],
            "notes": f"Welcome. Today I am presenting {startup_name}, {tagline}. Our primary goal is to address key challenges in our industry."
        },
        {
            "title": "The Problem",
            "bullets": [
                master_profile.get("problem", "A significant challenge faces our target users."),
                blueprint.get("problem_story", "Current methods are manual, inefficient, and slow.")
            ],
            "notes": "Let's start with the problem. Our customer segment deals with critical pain points daily. Current tools fail to solve this."
        },
        {
            "title": "The Solution",
            "bullets": [
                master_profile.get("solution", "Our product simplifies operations and unlocks new opportunities."),
                blueprint.get("solution_story", "We automate and optimize the workflow, creating direct value.")
            ],
            "notes": "Here is our solution. We provide a streamlined workflow that delivers tangible outcomes and significant time savings."
        },
        {
            "title": "Market Opportunity",
            "bullets": [
                blueprint.get("market_story", "The market opportunity is expanding rapidly."),
                market_opp_str,
                f"Sourced Data: {claims_strs[0]}"
            ],
            "notes": "Looking at the market. There is a verified and growing demand for our product. Here is our opportunity sizing."
        },
        {
            "title": "Key Capabilities",
            "bullets": [f"Feature: {f}" for f in master_profile.get("key_features", [])[:4]] or ["Core feature sets", "User-friendly panels", "Automated pipelines"],
            "notes": "Let's walk through what the product actually does. These key capabilities are designed to optimize daily operations."
        },
        {
            "title": "Business Model",
            "bullets": [
                blueprint.get("business_story", "Our pricing aligns directly with user value."),
                f"Revenue streams: {master_profile.get('business_model', 'Subscription SaaS')}"
            ],
            "notes": "How do we make money? Our business model is aligned with customer scale and features, creating a predictable recurring revenue flow."
        },
        {
            "title": "Roadmap & Traction",
            "bullets": [
                "Phase 1: Concept Validation and Market Research",
                "Phase 2: MVP Design and Branding Identity",
                "Phase 3: Landing Page Deployment and Early Sign-ups",
                "Phase 4: Scale and Feature Expansion"
            ],
            "notes": "Here is our roadmap. We have completed initial validation and are moving aggressively towards product launch and traction."
        },
        {
            "title": "Competitive Landscape",
            "bullets": [
                blueprint.get("competitive_story", "Differentiation lies in speed, simplicity, and premium design."),
                f"Identified Competitors: {competitors_str}"
            ],
            "notes": "Looking at competitors. We differentiate ourselves on customer service, pricing, and visual excellence."
        },
        {
            "title": "Go-to-Market Strategy",
            "bullets": [
                "Product-Led Growth (PLG) focusing on direct onboarding",
                "Content marketing highlighting key value propositions",
                "Strategic partnerships in target customer segments"
            ],
            "notes": "Our go-to-market strategy leverages organic growth loops and focused digital advertising to acquire early active users."
        },
        {
            "title": "Core Team",
            "bullets": [
                team_val
            ],
            "notes": "Let's look at the team. We have a solid engineering and design background to execute this product vision."
        },
        {
            "title": "Financial Projections",
            "bullets": [
                financials_val
            ],
            "notes": "Our financial projections are based on initial pricing models and targeted customer acquisition cost ratios."
        },
        {
            "title": "The Ask & Next Steps",
            "bullets": [
                blueprint.get("closing_story", "Partner with us to capture early market traction."),
                f"Action: {blueprint.get('call_to_action', 'Get Started')}"
            ],
            "notes": "We are seeking partners and early pilots to help validate our beta version. Thank you, and let's get started."
        }
    ]


def _generate_outline(master_profile: dict[str, Any]) -> list[dict[str, Any]]:
    prompt = (
        "You are an expert pitch deck copywriter and business analyst.\n"
        "Your task is to generate a list of exactly 12 slides for a startup pitch deck.\n"
        "Input Data: the Master Startup Profile containing industry, problem, solution, branding, narrative blueprint, and research findings.\n"
        "Format: Return a strict JSON object with a single 'slide_list' key containing an array of 12 objects. Each object must have:\n"
        "  - 'title': Slide title (max 50 chars)\n"
        "  - 'bullets': Array of 2 to 4 concise, high-impact bullet points (each max 120 chars)\n"
        "  - 'notes': Presenter speech script or notes (max 200 chars)\n"
        "\n"
        "Slide Order standard:\n"
        "1. Title / Intro (Startup name, tagline, hero message)\n"
        "2. The Problem (Pain points in customer segment)\n"
        "3. The Solution (Product value proposition)\n"
        "4. Market Opportunity (TAM/SAM/SOM opportunity size)\n"
        "5. Product Capabilities (Key feature sets)\n"
        "6. Business Model (Pricing and revenue streams)\n"
        "7. Roadmap & Traction (Milestones and next phases)\n"
        "8. Competitive Landscape (Competitors and unique differentiators)\n"
        "9. Go-to-Market Strategy (Marketing and user acquisition plan)\n"
        "10. Core Team (Team members or status)\n"
        "11. Financial Projections (Forecasts or status)\n"
        "12. The Ask / Next Steps (Call to action and contact information)\n"
        "\n"
        "Strict rules:\n"
        "1. No system/agent leakage: Never mention 'IdeaBee', 'agent', 'multi-agent', 'pipeline', 'prompt', 'database', etc.\n"
        "2. Team slide (Slide 10): If the team is not specified or is TBD, must strictly show only 'To Be Determined' as the bullet text. Do not invent names.\n"
        "3. Financials slide (Slide 11): If financials are not specified or empty, must strictly show only 'Not Provided' as the bullet text. Do not invent metrics.\n"
        "4. Competitor slide (Slide 8): If competitors are 'Competitor analysis unavailable', must show exactly that and explain how we validate differentiation.\n"
        "5. TAM/SAM/SOM (Slide 4): If market opportunity data is unavailable, use 'Market data unavailable'.\n"
        "6. Do not fabricate any other names, numbers, or details that are not in the profile.\n"
    )

    model = os.getenv("PITCH_MODEL", "google/gemma-3-27b-it:featherless-ai")
    fallback = _get_python_fallback_slides(master_profile)

    for provider in (
        lambda: openrouter_chat(
            model=model,
            messages=[
                {"role": "system", "content": prompt},
                {"role": "user", "content": json.dumps(master_profile, indent=2)},
            ],
            temperature=0.25,
        ),
        lambda: hf_chat(
            model=model,
            messages=[
                {"role": "system", "content": prompt},
                {"role": "user", "content": json.dumps(master_profile, indent=2)},
            ],
            temperature=0.25,
        ),
    ):
        try:
            response = provider()
            data = extract_json(response)
            slides = data.get("slide_list") or data.get("slides") or data
            if isinstance(slides, list) and len(slides) == 12:
                normalized: list[dict[str, Any]] = []
                for slide in slides:
                    if isinstance(slide, dict):
                        normalized.append(
                            {
                                "title": str(slide.get("title", "Slide")),
                                "bullets": [str(item) for item in slide.get("bullets", [])],
                                "notes": str(slide.get("notes", "")),
                            }
                        )
                if len(normalized) == 12:
                    return normalized
        except Exception:
            continue

    return fallback


def _build_pptx(
    slides: list[dict[str, Any]],
    output_path: Path,
    brand_identity: dict[str, Any]
) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    layout = prs.slide_layouts[6]  # Blank slide layout

    primary_color_hex = brand_identity.get("primary_color", "#0284c7")
    secondary_color_hex = brand_identity.get("secondary_color", "#64748b")
    font_name = brand_identity.get("typography", "Plus Jakarta Sans")
    if isinstance(font_name, dict):
        font_name = font_name.get("primary") or font_name.get("header") or font_name.get("body") or list(font_name.values())[0]
    font_name = str(font_name)
    if "," in font_name:
        font_name = font_name.split(",")[0].strip()
    if "&" in font_name:
        font_name = font_name.split("&")[0].strip()


    primary_rgb = hex_to_rgb(primary_color_hex)
    secondary_rgb = hex_to_rgb(secondary_color_hex)
    white_rgb = RGBColor(255, 255, 255)
    dark_bg_rgb = RGBColor(11, 15, 25) # Sleek deep dark color #0b0f19
    light_grey_rgb = RGBColor(203, 213, 225) # #cbd5e1

    for index, slide_data in enumerate(slides):
        slide = prs.slides.add_slide(layout)

        # Set slide background
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = dark_bg_rgb

        # Draw decorative left accent line
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(0.5), Inches(0.08), Inches(6.5))
        shape.fill.solid()
        shape.fill.fore_color.rgb = primary_rgb
        shape.line.color.rgb = primary_rgb

        if index == 0:
            # Title Slide Layout
            textbox = slide.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.0), Inches(4.5))
            tf = textbox.text_frame
            tf.word_wrap = True
            tf.clear()

            # Startup Name
            p_name = tf.paragraphs[0]
            p_name.text = slide_data["title"]
            p_name.font.size = Pt(56)
            p_name.font.bold = True
            p_name.font.color.rgb = white_rgb
            p_name.font.name = font_name

            # Tagline
            bullets = slide_data.get("bullets", [])
            if bullets:
                p_tag = tf.add_paragraph()
                p_tag.text = bullets[0]
                p_tag.font.size = Pt(24)
                p_tag.font.color.rgb = primary_rgb
                p_tag.font.name = font_name
                p_tag.space_before = Pt(14)

                for b in bullets[1:]:
                    p_sub = tf.add_paragraph()
                    p_sub.text = b
                    p_sub.font.size = Pt(18)
                    p_sub.font.color.rgb = light_grey_rgb
                    p_sub.font.name = font_name
                    p_sub.space_before = Pt(8)
        else:
            # Content Slide Layout
            # Slide Title
            title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.733), Inches(0.8))
            tf_title = title_box.text_frame
            tf_title.word_wrap = True
            p_title = tf_title.paragraphs[0]
            p_title.text = slide_data["title"]
            p_title.font.size = Pt(32)
            p_title.font.bold = True
            p_title.font.color.rgb = white_rgb
            p_title.font.name = font_name

            # Body TextBox
            body_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(11.733), Inches(4.6))
            tf_body = body_box.text_frame
            tf_body.word_wrap = True
            tf_body.clear()

            bullets = slide_data.get("bullets", [])
            for bullet_index, bullet_text in enumerate(bullets):
                p_bullet = tf_body.paragraphs[0] if bullet_index == 0 else tf_body.add_paragraph()
                p_bullet.text = f"•  {bullet_text}"
                p_bullet.font.size = Pt(20)
                p_bullet.font.color.rgb = light_grey_rgb
                p_bullet.font.name = font_name
                p_bullet.space_before = Pt(12)

        # Presenter notes
        if slide_data.get("notes"):
            slide.notes_slide.notes_text_frame.text = slide_data.get("notes", "")

        # Footer & Slide Number
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.733), Inches(0.35))
        tf_footer = footer_box.text_frame
        p_foot = tf_footer.paragraphs[0]
        p_foot.text = f"Confidential  |  Widescreen Investor Presentation  |  Slide {index + 1}"
        p_foot.font.size = Pt(10)
        p_foot.font.color.rgb = secondary_rgb
        p_foot.font.name = font_name

    prs.save(str(output_path))


def _wrap_text_reportlab(text: str, max_chars: int) -> list[str]:
    words = text.split(" ")
    lines = []
    current_line = []
    current_len = 0
    for word in words:
        if current_len + len(word) + 1 > max_chars:
            lines.append(" ".join(current_line))
            current_line = [word]
            current_len = len(word)
        else:
            current_line.append(word)
            current_len += len(word) + 1
    if current_line:
        lines.append(" ".join(current_line))
    return lines


def _export_pdf_fallback(
    slides: list[dict[str, Any]],
    pdf_path: Path,
    brand_identity: dict[str, Any],
    startup_name: str
) -> None:
    from reportlab.lib.colors import HexColor
    from reportlab.pdfgen import canvas

    width, height = 960, 540
    c = canvas.Canvas(str(pdf_path), pagesize=(width, height))

    primary_hex = brand_identity.get("primary_color", "#0284c7")
    secondary_hex = brand_identity.get("secondary_color", "#64748b")

    for index, slide in enumerate(slides):
        # Slate background
        c.setFillColor(HexColor("#0b0f19"))
        c.rect(0, 0, width, height, fill=1, stroke=0)

        # Left bar accent
        c.setFillColor(HexColor(primary_hex))
        c.rect(30, 40, 6, height - 80, fill=1, stroke=0)

        if index == 0:
            # Title slide
            c.setFont("Helvetica-Bold", 42)
            c.setFillColor(HexColor("#ffffff"))
            c.drawString(60, height - 180, slide.get("title", "Startup"))

            c.setFont("Helvetica-Bold", 20)
            c.setFillColor(HexColor(primary_hex))
            bullets = slide.get("bullets", [])
            y = height - 240
            if bullets:
                c.drawString(60, y, bullets[0])
                y -= 40
                c.setFont("Helvetica", 14)
                c.setFillColor(HexColor("#cbd5e1"))
                for b in bullets[1:]:
                    wrapped_lines = _wrap_text_reportlab(b, 80)
                    for wl in wrapped_lines:
                        c.drawString(60, y, wl)
                        y -= 25
        else:
            # Content slides
            # Slide Title
            c.setFont("Helvetica-Bold", 26)
            c.setFillColor(HexColor("#ffffff"))
            c.drawString(60, height - 80, slide.get("title", "Slide"))

            # Bullets
            c.setFont("Helvetica", 15)
            c.setFillColor(HexColor("#cbd5e1"))
            y = height - 140
            for bullet in slide.get("bullets", []):
                wrapped_lines = _wrap_text_reportlab(bullet, 85)
                for i, line in enumerate(wrapped_lines):
                    bullet_marker = "•  " if i == 0 else "   "
                    c.drawString(80, y, f"{bullet_marker}{line}")
                    y -= 25
                    if y < 90:
                        break
                y -= 10
                if y < 90:
                    break

        # Presenter notes
        notes = slide.get("notes", "")
        if notes:
            c.setFont("Helvetica-Oblique", 10)
            c.setFillColor(HexColor("#94a3b8"))
            c.drawString(60, 50, f"Notes: {notes[:140]}")

        # Footer
        c.setFont("Helvetica", 9)
        c.setFillColor(HexColor(secondary_hex))
        c.drawRightString(width - 60, 25, f"Confidential  |  {startup_name} Widescreen Deck  |  Slide {index + 1}")

        c.showPage()

    c.save()


def build_pitch_deck(master_profile: dict[str, Any], output_dir: str = "outputs") -> dict[str, Any]:
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)

    pptx_path = output_path / "startup_pitch.pptx"
    pdf_path = output_path / "startup_pitch.pdf"

    # 1. Generate Slide List content outline
    slides = _generate_outline(master_profile)

    # 2. Build Widescreen PowerPoint Presentation programmatically
    brand_identity = master_profile.get("brand_identity", {})
    _build_pptx(slides, pptx_path, brand_identity)

    # 3. Save slides to JSON
    slides_json = output_path / "startup_pitch_slides.json"
    slides_json.write_text(json.dumps(slides, indent=2), encoding="utf-8")

    # 4. Generate widescreen direct PDF slide deck
    startup_name = master_profile.get("startup_name", "Startup")
    _export_pdf_fallback(slides, pdf_path, brand_identity, startup_name)

    return {
        "status": "generated",
        "slides": slides,
        "slides_json": str(slides_json),
        "pptx": str(pptx_path),
        "pdf": str(pdf_path),
        "brand_identity": brand_identity,
        "master_profile": master_profile,
    }
